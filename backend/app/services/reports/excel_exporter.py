import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from app.services.reports.data import ReportData

_HEADER_FILL = PatternFill(start_color="1F2733", end_color="1F2733", fill_type="solid")
_HEADER_FONT = Font(color="22D3EE", bold=True)
_TITLE_FONT = Font(size=14, bold=True)


def _write_table(ws: Worksheet, start_row: int, headers: list[str], rows: list[list], title: str | None = None):
    row = start_row
    if title:
        ws.cell(row=row, column=1, value=title).font = Font(bold=True)
        row += 1

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    row += 1

    for data_row in rows:
        for col, value in enumerate(data_row, start=1):
            ws.cell(row=row, column=col, value=value)
        row += 1

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col)].width = 20

    return row + 1  # next free row, with a blank line gap


def export_excel(report: ReportData) -> bytes:
    wb = Workbook()

    summary_ws = wb.active
    summary_ws.title = "Summary"
    summary_ws["A1"] = report.period_label
    summary_ws["A1"].font = _TITLE_FONT
    summary_ws["A2"] = f"Generated: {report.generated_at.isoformat()}"
    summary_ws["A2"].alignment = Alignment(horizontal="left")

    row = _write_table(
        summary_ws,
        4,
        ["Metric", "Value"],
        [["Total Attacks", report.total_attacks], ["Unique Attacker IPs", report.unique_attacker_ips]],
    )
    _write_table(
        summary_ws,
        row,
        ["Severity", "Count"],
        [[r["severity"], r["attack_count"]] for r in report.severity_breakdown],
        title="Severity Breakdown",
    )

    ips_ws = wb.create_sheet("Top Attacker IPs")
    _write_table(
        ips_ws, 1, ["Source IP", "Attack Count"], [[r["source_ip"], r["attack_count"]] for r in report.top_attacker_ips]
    )

    countries_ws = wb.create_sheet("Top Countries")
    _write_table(
        countries_ws,
        1,
        ["Country Code", "Attack Count"],
        [[r["country_code"], r["attack_count"]] for r in report.top_countries],
    )

    creds_ws = wb.create_sheet("Top Credentials")
    _write_table(
        creds_ws,
        1,
        ["Username", "Password", "Attempt Count"],
        [[r["username"], r["password"], r["attempt_count"]] for r in report.top_credentials],
    )

    sessions_ws = wb.create_sheet("Sessions")
    _write_table(
        sessions_ws,
        1,
        ["Source IP", "Protocol", "Country", "Severity", "Risk Score", "Started At", "Duration (s)"],
        [
            [
                s.source_ip,
                s.protocol,
                s.country_code or "",
                s.severity.value if s.severity else "unclassified",
                s.risk_score if s.risk_score is not None else "",
                s.started_at.isoformat(),
                s.duration_seconds if s.duration_seconds is not None else "",
            ]
            for s in report.sessions
        ],
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
